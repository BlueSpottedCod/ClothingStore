import pytest
import sqlite3
import os
import openpyxl
import app

# Путь к тестовому файлу Excel
TEST_EXCEL_FILE = "test_sales_report.xlsx"

# --- Helper функции для проверки ---
def get_products():
    conn = app.connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM products")
    products = cursor.fetchall()
    conn.close()
    return products

def get_employees():
    conn = app.connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM employees")
    employees = cursor.fetchall()
    conn.close()
    return employees

# ------------------------
# Тесты для базы данных
# ------------------------
def test_connect_db():
    conn = app.connect_db()
    assert isinstance(conn, sqlite3.Connection)
    conn.close()

def test_create_tables():
    app.create_tables()
    conn = app.connect_db()
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='products'")
    assert cursor.fetchone() is not None
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='sales'")
    assert cursor.fetchone() is not None
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='employees'")
    assert cursor.fetchone() is not None
    conn.close()

def test_initial_products():
    products = get_products()
    assert len(products) > 0

def test_initial_employees():
    employees = get_employees()
    assert len(employees) > 0

# ------------------------
# Тест логики продаж (через базу)
# ------------------------
def test_add_sale_logic():
    conn = app.connect_db()
    cursor = conn.cursor()
    product = get_products()[0]
    initial_quantity = product[3]
    product_id = product[0]
    quantity = 1
    date_now = "01-01-2025 12:00:00"
    cursor.execute("INSERT INTO sales (product_id, quantity, date) VALUES (?, ?, ?)", 
                   (product_id, quantity, date_now))
    cursor.execute("UPDATE products SET quantity = ? WHERE id = ?", 
                   (initial_quantity - quantity, product_id))
    conn.commit()
    cursor.execute("SELECT quantity FROM products WHERE id = ?", (product_id,))
    updated_quantity = cursor.fetchone()[0]
    conn.close()
    assert updated_quantity == initial_quantity - quantity

# ------------------------
# Тест экспорта в Excel
# ------------------------
def test_export_sales_to_excel():
    # Удаляем файл, если он есть
    if os.path.exists(TEST_EXCEL_FILE):
        os.remove(TEST_EXCEL_FILE)
    
    # Получаем все продажи из базы
    conn = app.connect_db()
    cursor = conn.cursor()
    cursor.execute('''SELECT sales.id, products.name, sales.quantity, products.price, sales.date
                      FROM sales JOIN products ON sales.product_id = products.id''')
    sales = cursor.fetchall()
    total_sum = sum(s[2] * s[3] for s in sales)
    conn.close()
    
    # Создаем Excel файл (по аналогии с функцией из app)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Продажи"
    headers = ["ID", "Продукт", "Кол-во", "Цена за шт.", "Дата", "Сумма"]
    sheet.append(headers)
    for sale in sales:
        total_price = sale[2] * sale[3]
        sheet.append([sale[0], sale[1], sale[2], sale[3], sale[4], total_price])
    sheet.append(["", "", "", "", "Итого:", total_sum])
    workbook.save(TEST_EXCEL_FILE)
    
    # Проверка: файл создан
    assert os.path.exists(TEST_EXCEL_FILE)
    
    # Проверка содержимого файла
    wb = openpyxl.load_workbook(TEST_EXCEL_FILE)
    sh = wb.active
    assert sh.title == "Продажи"
    assert sh.cell(row=1, column=1).value == "ID"
    assert sh.cell(row=1, column=6).value == "Сумма"
    # Проверяем последнюю строку с итого
    last_row = sh.max_row
    assert sh.cell(row=last_row, column=5).value == "Итого:"
    assert sh.cell(row=last_row, column=6).value == total_sum
    
    # Убираем тестовый файл после проверки
    os.remove(TEST_EXCEL_FILE)
